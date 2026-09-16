import re
import unicodedata
import json
from pathlib import Path

from app.services.analysis import analisar_colunas, analisar_csv, ler_csv
from app.services.ai import (
    preparar_contexto_ia,
    criar_prompt_analise,
    consultar_ia
)

TABLE_EXTENSIONS = {".csv", ".xls", ".xlsx"}
TEXT_EXTENSIONS = {".pdf", ".docx", ".txt"}
SUPPORTED_EXTENSIONS = TABLE_EXTENSIONS | TEXT_EXTENSIONS


def _read_text_file(file_path: str) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin1"):
        try:
            return Path(file_path).read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    raise ValueError("Não foi possível ler o texto deste arquivo.")


def _read_pdf(file_path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ValueError("O suporte a PDF não está instalado no backend.") from error

    reader = PdfReader(file_path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def _read_docx(file_path: str) -> str:
    try:
        from docx import Document
    except ImportError as error:
        raise ValueError("O suporte a DOCX não está instalado no backend.") from error

    document = Document(file_path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [
        " | ".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + tables).strip()

def _ler_dimensao_tabela(file_path: str) -> dict:
    df = ler_csv(file_path)

    return {
        "total_registros": int(len(df)),
        "total_colunas": int(len(df.columns)),
        "colunas": list(df.columns)
    }

def read_document(file_path: str) -> dict:
    if file_path.startswith("database:"):
        import pandas as pd

        config = json.loads(file_path.removeprefix("database:"))
        rows = config.get("rows", [])
        frame = pd.DataFrame(rows)
        return {
            "kind": "table",
            "analysis": {
                "total_registros": int(len(frame)),
                "total_colunas": int(len(frame.columns)),
                "colunas": [str(column) for column in frame.columns],
                "analise_colunas": analisar_colunas(frame),
                "insights": {}
            },
            "text": ""
        }

    extension = Path(file_path).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Formato não suportado. Use PDF, DOCX, TXT, CSV, XLS ou XLSX."
        )

    if extension in TABLE_EXTENSIONS:
        return {
            "kind": "table",
            "analysis": analisar_csv(file_path),
            "text": ""
        }

    if extension == ".pdf":
        text = _read_pdf(file_path)
    elif extension == ".docx":
        text = _read_docx(file_path)
    else:
        text = _read_text_file(file_path)

    if not text.strip():
        raise ValueError("Não foi possível encontrar texto no arquivo.")

    return {
        "kind": "text",
        "analysis": {
            "total_caracteres": len(text),
            "total_palavras": len(re.findall(r"\S+", text)),
            "total_paragrafos": len([line for line in text.splitlines() if line.strip()])
        },
        "text": text
    }


def _table_answer(question: str, analysis: dict) -> str:
    lower_question = _normalize(question)
    total_records = analysis.get("total_registros", 0)
    total_columns = analysis.get("total_colunas", 0)
    columns = analysis.get("colunas", [])
    insights = analysis.get("insights", {})
    column_analysis = analysis.get("analise_colunas", {})

    mentioned_column = next(
        (
            column for column in columns
            if _normalize(str(column)) in lower_question
        ),
        None
    )

    if mentioned_column:
        details = column_analysis.get(mentioned_column, {})
        answer = [
            f"Análise da coluna {mentioned_column}",
            f"Tipo: {details.get('tipo', 'não identificado')}.",
            f"Valores preenchidos: {details.get('valores_preenchidos', 0)}.",
            f"Valores nulos: {details.get('valores_nulos', 0)}.",
            f"Valores únicos: {details.get('valores_unicos', 0)}."
        ]
        statistics = details.get("estatisticas")
        if statistics:
            answer.append(
                "Estatísticas: "
                f"mínimo {statistics['minimo']}, "
                f"máximo {statistics['maximo']}, "
                f"média {statistics['media']} e "
                f"mediana {statistics['mediana']}."
            )
        return "\n\n".join(answer)

    if (
        any(term in lower_question for term in ("registro", "registros", "linha", "linhas"))
        and any(term in lower_question for term in ("coluna", "colunas", "campo", "campos", "variavel", "variaveis"))
    ):
        return (
            f"O arquivo possui {total_records} registros "
            f"e {total_columns} colunas."
        )

    if any(term in lower_question for term in ("linha", "registro", "quantidade", "total")):
        return "\n\n".join([
            "Dimensão do arquivo",
            f"Foram encontrados {total_records} registros e {total_columns} colunas.",
            "Esse número foi calculado diretamente a partir da tabela enviada."
        ])

    highlights = []
    for group in ("gerais", "relacoes", "comparativos", "qualidade_dados"):
        values = insights.get(group, [])
        if isinstance(values, list):
            highlights.extend(str(value) for value in values[:3])

    asks_for_summary = _is_summary_question(lower_question)
    if highlights or asks_for_summary:
        return "\n\n".join([
            "Resumo da análise",
            f"O arquivo contém {total_records} registros distribuídos em {total_columns} colunas.",
            "Principais achados:\n- " + "\n- ".join(highlights[:4])
            if highlights else "Não foram calculados achados categóricos para este arquivo."
        ])

    return "\n\n".join([
        "Resultado da análise",
        f"O arquivo contém {total_records} registros e {total_columns} colunas.",
        "Não foi identificado um achado específico para essa pergunta."
    ])


def _text_answer(question: str, text: str) -> str:
    normalized_question = _normalize(question)
    keywords = {
        word for word in normalized_question.split()
        if len(word) >= 4 and word not in {
            "qual", "quais", "como", "onde", "quando", "sobre", "resumindo",
            "resuma", "resumo", "assunto", "tema", "principal", "isso",
            "esse", "esta", "para", "arquivo", "documento"
        }
    }
    paragraphs = _text_units(text)

    if _is_summary_question(normalized_question):
        return _text_summary(text, paragraphs)

    ranked = sorted(
        (
            (sum(keyword in _normalize(paragraph) for keyword in keywords), index, paragraph)
            for index, paragraph in enumerate(paragraphs)
        ),
        key=lambda item: (item[0], -item[1]),
        reverse=True
    )
    matches = [paragraph for score, _, paragraph in ranked if score > 0][:3]

    if matches:
        return "\n\n".join([
            "Resposta baseada no documento",
            "Encontrei estes trechos relacionados à sua pergunta:",
            "- " + "\n- ".join(matches),
            "A resposta foi construída apenas com o texto extraído do arquivo."
        ])

    excerpt = " ".join(text.split())[:700]
    return "\n\n".join([
        "Informação não localizada",
        "Não encontrei termos diretamente relacionados à pergunta no documento.",
        f"Trecho inicial disponível: {excerpt}",
        "Tente perguntar usando palavras que aparecem no arquivo."
    ])


def _normalize(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    without_accents = "".join(
        character for character in normalized
        if not unicodedata.combining(character)
    )
    return re.sub(r"[^\w\s]", " ", without_accents.lower())

def _selecionar_trechos_relevantes(question: str, text: str) -> list[str]:
    keywords = {
        word
        for word in _normalize(question).split()
        if len(word) >= 4
        and word not in {
            "qual", "quais", "como", "onde", "quando",
            "sobre", "esse", "esta", "para",
            "arquivo", "documento"
        }
    }

    paragraphs = [
        " ".join(paragraph.split())
        for paragraph in re.split(
            r"\n\s*\n|(?<=[.!?])\s+",
            text
        )
        if paragraph.strip()
    ]

    ranked = sorted(
        (
            (
                sum(
                    keyword in _normalize(paragraph)
                    for keyword in keywords
                ),
                index,
                paragraph
            )
            for index, paragraph in enumerate(paragraphs)
        ),
        key=lambda item: (item[0], -item[1]),
        reverse=True
    )

    return [
        paragraph
        for score, _, paragraph in ranked
        if score > 0
    ][:3]

def _is_summary_question(normalized_question: str) -> bool:
    return any(
        term in normalized_question
        for term in (
            "resum", "assunto", "tema", "sobre o que", "do que trata",
            "principais pontos", "em poucas palavras"
        )
    )


def _text_units(text: str) -> list[str]:
    normalized = " ".join(text.split())
    numbered_units = re.split(r"(?=\b\d+\.\s+)", normalized)
    if len(numbered_units) > 1:
        units = numbered_units
    else:
        units = re.split(r"(?<=[.!?])\s+", normalized)

    return [unit.strip(" -") for unit in units if unit.strip()]


def _text_summary(text: str, units: list[str]) -> str:
    normalized_text = _normalize(text)
    topic = "o conteúdo apresentado no documento"
    topic_signals = [
        (("treino", "exercicio", "serie", "reps", "aquecimento"), "um plano de treino e exercícios físicos"),
        (("aluno", "escola", "professor", "educacao"), "informações educacionais"),
        (("venda", "produto", "cliente", "faturamento"), "dados comerciais e de clientes"),
        (("receita", "ingrediente", "preparo"), "uma receita ou instruções de preparo")
    ]
    for signals, description in topic_signals:
        if sum(signal in normalized_text for signal in signals) >= 2:
            topic = description
            break

    points = []
    for unit in units:
        compact = unit[:220].strip()
        if compact and compact not in points:
            points.append(compact)
        if len(points) == 4:
            break

    return "\n\n".join([
        "Resumo do documento",
        f"O documento trata de {topic}.",
        "Pontos identificados no conteúdo:\n- " + "\n- ".join(points),
        "Resumo construído exclusivamente a partir do texto extraído do arquivo."
    ])


def answer_question(file_path: str, source_name: str, question: str) -> dict:
    normalized_question = _normalize(question)

    pergunta_dimensao = any(
        termo in normalized_question
        for termo in (
            "registro",
            "registros",
            "linha",
            "linhas",
            "quantidade de registros",
            "total de registros",
            "quantas linhas"
        )
    )

    pergunta_colunas = any(
        termo in normalized_question
        for termo in (
            "coluna",
            "colunas",
            "campo",
            "campos",
            "variavel",
            "variaveis"
        )
    )

    extension = Path(file_path).suffix.lower()

    try:
        # Perguntas simples sobre quantidade de registros/colunas
        # não precisam executar a análise completa do CSV.
        if extension in TABLE_EXTENSIONS and (
            pergunta_dimensao or pergunta_colunas
        ):
            analysis = _ler_dimensao_tabela(file_path)

            answer = _table_answer(
                question,
                analysis
            )

            return {
                "source_name": source_name,
                "question": question,
                "answer": answer,
                "kind": "table",
                "analysis": analysis
            }

        # Demais perguntas continuam usando o fluxo completo.
        document = read_document(file_path)

        if document["kind"] == "table":

            contexto = preparar_contexto_ia(
                document["analysis"],
                question
            )

            prompt = criar_prompt_analise(
                contexto,
                question
            )

            resposta_ia = consultar_ia(prompt)
            answer = resposta_ia

        else:
            trechos = _selecionar_trechos_relevantes(
                question,
                document["text"]
            )

            prompt = f"""
Você é um assistente de análise de documentos.

Responda sempre em português do Brasil.

Responda diretamente à pergunta do usuário.

Use somente as informações presentes nos trechos fornecidos.

Não invente informações.

Não mostre seu raciocínio.

Se a informação não estiver nos trechos, diga que ela não foi localizada.

PERGUNTA DO USUÁRIO:

{question}

TRECHOS RELEVANTES DO DOCUMENTO:

{chr(10).join("- " + trecho for trecho in trechos)}

Responda de forma clara e objetiva.
"""

            resposta_ia = consultar_ia(prompt)
            answer = resposta_ia

    except Exception:
        document = read_document(file_path)

        if document["kind"] == "table":
            answer = _table_answer(
                question,
                document["analysis"]
            )
        else:
            answer = _text_answer(
                question,
                document["text"]
            )

    return {
        "source_name": source_name,
        "question": question,
        "answer": answer,
        "kind": document["kind"],
        "analysis": document["analysis"]
    }